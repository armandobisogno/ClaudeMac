# -*- coding: utf-8 -*-
import re, html
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_ORIENT

SRC = "/private/tmp/claude-501/-Users-macbook-Library-CloudStorage-Dropbox-ClaudeWS/43048cf3-9698-4d82-831f-44ad0b2ce3d8/scratchpad/contr_acc"
ROMAN = ["I", "II", "III"]
N_BOOKS = 3

# ---------- lettura file ----------
def get_body(path):
    d = open(path, encoding="cp1252", errors="replace").read()
    lo = d.lower()
    a = lo.find("<body")
    a = d.find(">", a) + 1
    b = lo.find("</body>")
    return d[a:b]

BLOCK = re.compile(r"<(h3|h4|h5|p)\b([^>]*)>(.*?)</\1>", re.I | re.S)
KEY   = re.compile(r'<a\s+name="A_(\d{3})_(\d{3})_(\d{3})"', re.I)

def clean_inline(s):
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.I)
    s = re.sub(r"</?(?:a|u|sup|sub|font|span|h[1-6]|p|div)\b[^>]*>", "", s, flags=re.I)
    return s

def to_runs(s):
    s = clean_inline(s)
    parts = re.split(r"(<i>|</i>|<b>|</b>)", s, flags=re.I)
    bold = ital = 0
    out = []
    for p in parts:
        if not p:
            continue
        pl = p.lower()
        if pl == "<i>": ital += 1; continue
        if pl == "</i>": ital = max(0, ital - 1); continue
        if pl == "<b>": bold += 1; continue
        if pl == "</b>": bold = max(0, bold - 1); continue
        txt = re.sub(r"<[^>]+>", "", p)
        txt = html.unescape(txt).replace("\xa0", " ")
        txt = re.sub(r"[ \t\r\n]+", " ", txt)
        if txt:
            out.append(("t", txt, bool(bold), bool(ital)))
    ti = [i for i, r in enumerate(out) if r[0] == "t"]
    if ti:
        f = ti[0]; out[f] = ("t", out[f][1].lstrip(), out[f][2], out[f][3])
        l = ti[-1]; out[l] = ("t", out[l][1].rstrip(), out[l][2], out[l][3])
        out = [r for r in out if r[0] != "t" or r[1]]
    return out

# ---------- intestazione di libro ----------
def book_header(body, lang):
    """Ritorna il sottotitolo/argomento del libro (riga in maiuscolo dopo il titolo LIBRO N)."""
    if lang == "lat":
        m = re.search(r"<h3\b[^>]*>.*?</h3>\s*<b>\s*<p\b[^>]*>(.*?)</p>", body, re.I | re.S)
    else:
        m = re.search(r"<h4\b[^>]*>.*?</h4>\s*<h3\b[^>]*>(.*?)</h3>", body, re.I | re.S)
    if not m:
        return "", body
    t = re.sub(r"<br\s*/?>", " ", m.group(1), flags=re.I)
    t = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", t))).strip()
    return t, body[m.end():]

# ---------- corpo del libro ----------
def parse_body(body):
    """items: ('h', livello, runs) oppure ('p', (cap,par) | None, runs)."""
    raw = [(t.lower(), attrs, inner) for t, attrs, inner in BLOCK.findall(body)]
    items = []
    for idx, (tag, attrs, inner) in enumerate(raw):
        text_only = html.unescape(re.sub(r"<[^>]+>", "", inner)).replace("\xa0", " ").strip()
        if tag in ("h4", "h5"):
            if not text_only:
                continue
            items.append(("h", 4 if tag == "h4" else 5, to_runs(inner)))
            continue
        if tag == "h3":
            continue  # non dovrebbe comparire nel corpo (solo nell'intestazione)
        # tag == 'p'
        km = KEY.search(inner)
        if km:
            items.append(("p", (km.group(2), km.group(3)), to_runs(inner)))
            continue
        is_center = bool(re.search(r'align\s*=\s*"?center', attrs, re.I))
        if is_center:
            if not text_only:
                continue
            nxt = raw[idx + 1][0] if idx + 1 < len(raw) else None
            if nxt == "h5":
                items.append(("h", 4, to_runs(inner)))
                continue
            # altrimenti: citazione centrata incorporata nel paragrafo (es. versi di Virgilio)
        if not text_only:
            continue
        items.append(("p", None, to_runs(inner)))
    return items

def parse_book(path, lang):
    body = get_body(path)
    theme, rest = book_header(body, lang)
    return theme, parse_body(rest)

def grouped(items):
    groups = {}
    order = []
    pending = []
    for it in items:
        if it[0] == "h":
            pending.append(it)
        else:
            key = it[1]
            if key is None:
                if order:
                    groups[order[-1]]["runs"].append(("t", " ", False, False))
                    groups[order[-1]]["runs"].extend(it[2])
                continue
            if key not in groups:
                groups[key] = {"heads": [], "runs": []}
                order.append(key)
            groups[key]["heads"].extend(pending)
            groups[key]["runs"].extend(it[2])
            pending = []
    return order, groups

# ---------- citazioni bibliche/classiche incorporate nel testo ----------
# Sul sito le citazioni (bibliche o di autori classici) sono scritte fra parentesi
# direttamente nel corpo del testo, identiche nella versione latina e in quella
# italiana (sono un'aggiunta editoriale moderna, non parte della prosa di
# Agostino). Le estraiamo e le trasformiamo in note a fondo di libro.
CITATION_RE = re.compile(r'\(((?:Cf\.\s*)?\d*\s*[A-ZÀ-Ú][^()]*?\d[^()]*?)\)')

def extract_citations(runs):
    """Da una lista di run ('t', testo, bold, italic) estrae le citazioni fra
    parentesi (anche a cavallo di run diversi per via del corsivo) e le
    sostituisce con un marcatore ('n', indice_locale_1based).
    Ritorna (nuovi_runs, [testo_citazione, ...])."""
    offsets = []
    pos = 0
    for ridx, r in enumerate(runs):
        text = r[1]
        offsets.append((pos, pos + len(text), ridx))
        pos += len(text)
    flat = "".join(r[1] for r in runs)

    matches = list(CITATION_RE.finditer(flat))
    if not matches:
        return runs, []

    def run_slice(g_start, g_end):
        for (rs, re_, ridx) in offsets:
            if re_ <= g_start or rs >= g_end:
                continue
            s = max(rs, g_start) - rs
            e = min(re_, g_end) - rs
            seg = runs[ridx][1][s:e]
            if seg:
                yield (seg, runs[ridx][2], runs[ridx][3])

    citations = []
    new_runs = []
    cur = 0
    for m in matches:
        gs, ge = m.start(), m.end()
        if gs > cur:
            for seg, b, i in run_slice(cur, gs):
                new_runs.append(("t", seg, b, i))
        citations.append(re.sub(r"\s+", " ", m.group(1)).strip())
        new_runs.append(("n", len(citations)))
        cur = ge
    if cur < len(flat):
        for seg, b, i in run_slice(cur, len(flat)):
            new_runs.append(("t", seg, b, i))
    return new_runs, citations

def apply_footnotes(lgrp, igrp, counter, footnotes):
    """Estrae le citazioni da entrambe le colonne per una stessa chiave
    (paragrafo), assegna un numero di nota condiviso in ordine di comparsa
    e rimappa i marcatori locali su quel numero. counter è una lista [int]
    usata come contatore mutabile fra le chiamate."""
    l_runs, l_cits = extract_citations(lgrp["runs"])
    i_runs, i_cits = extract_citations(igrp["runs"])
    n = max(len(l_cits), len(i_cits))
    numbers = []
    for idx in range(n):
        counter[0] += 1
        text = i_cits[idx] if idx < len(i_cits) else l_cits[idx]
        footnotes[counter[0]] = text
        numbers.append(counter[0])

    def remap(rs):
        return [("n", str(numbers[r[1] - 1])) if r[0] == "n" else r for r in rs]

    lgrp["runs"] = remap(l_runs)
    igrp["runs"] = remap(i_runs)

# ---------- docx helpers ----------
def add_runs(par, runs, base_size=None):
    for r in runs:
        if r[0] == "n":
            run = par.add_run(r[1])
            run.font.superscript = True
            if base_size: run.font.size = base_size
            continue
        _, txt, b, i = r
        run = par.add_run(txt)
        run.bold = b
        run.italic = i
        if base_size: run.font.size = base_size

def write_note_list(doc, heading, footnotes):
    if not footnotes:
        return
    doc.add_paragraph()
    np = doc.add_paragraph()
    nr = np.add_run(heading); nr.bold = True; nr.font.size = Pt(11)
    for num in sorted(footnotes):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        rn = p.add_run(f"{num}. "); rn.bold = True; rn.font.size = Pt(8.5)
        rt = p.add_run(footnotes[num]); rt.font.size = Pt(8.5)

def set_col_widths(table, w):
    for row in table.rows:
        for cell in row.cells:
            cell.width = w

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)

FONT = "Garamond"

def force_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)

# ---------- build ----------
doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)
st.paragraph_format.line_spacing = 1.15
for sname in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    try:
        stx = doc.styles[sname]
        stx.font.name = FONT
        stx.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except KeyError:
        pass

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1.6)

CELL_W = Cm(13.3)

title = doc.add_heading("Agostino d’Ippona — Contro gli Accademici", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Testo latino e traduzione italiana a fronte")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

FONTE = (
    "Fonte di questa versione. Il testo latino e la traduzione italiana sono ripresi "
    "dall’edizione digitale pubblicata sul sito www.augustinus.it:\n"
    "· testo latino: https://www.augustinus.it/latino/contr_acc/ (Libri I–III);\n"
    "· traduzione italiana: https://www.augustinus.it/italiano/contr_acc/ (Libri I–III).\n"
    "Consultazione: 22 settembre 2026. La numerazione di capitoli e paragrafi segue quella "
    "dell’edizione. Il sito non prevede, per quest’opera, un apparato di note separato: le "
    "citazioni bibliche e classiche sono scritte in originale fra parentesi nel corpo del testo, "
    "identiche nella versione latina e in quella italiana. In questa trascrizione sono state "
    "estratte e riportate come note a fondo di ciascun libro, con richiamo numerato in entrambe "
    "le colonne."
)
for line in FONTE.split("\n"):
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: r.font.size = Pt(8.5); r.italic = True

for bi in range(1, N_BOOKS + 1):
    ii = "%02d" % bi
    lat_theme, lat_items = parse_book(f"{SRC}/latino_{ii}_libro.html", "lat")
    ita_theme, ita_items = parse_book(f"{SRC}/italiano_{ii}_libro.html", "ita")

    doc.add_page_break()
    h = doc.add_heading(f"LIBRO {ROMAN[bi-1]}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp = doc.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = hp.add_run(lat_theme); r1.italic = True; r1.bold = True
    if ita_theme:
        hp2 = doc.add_paragraph(); hp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = hp2.add_run(ita_theme); r2.bold = True
        r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    lo, lg = grouped(lat_items)
    io, ig = grouped(ita_items)

    keys = list(lo)
    for k in io:
        if k not in lg:
            keys.append(k)

    footnotes = {}
    counter = [0]
    for k in keys:
        lgrp = lg.get(k)
        igrp = ig.get(k)
        if lgrp is None: lgrp = lg[k] = {"heads": [], "runs": []}
        if igrp is None: igrp = ig[k] = {"heads": [], "runs": []}
        apply_footnotes(lgrp, igrp, counter, footnotes)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for c, lbl in zip(hdr, ("LATINO", "ITALIANO")):
        c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(lbl); rr.bold = True
        shade(c, "E8E8E8")

    for k in keys:
        lgrp = lg.get(k, {"heads": [], "runs": []})
        igrp = ig.get(k, {"heads": [], "runs": []})
        lh, ih = lgrp["heads"], igrp["heads"]
        nh = max(len(lh), len(ih))
        for hidx in range(nh):
            row = table.add_row().cells
            for cell, hs in ((row[0], lh), (row[1], ih)):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if hidx < len(hs):
                    lvl = hs[hidx][1]
                    add_runs(p, hs[hidx][2])
                    for rn in p.runs:
                        rn.bold = True
                        rn.italic = (lvl == 5)
                        rn.font.size = Pt(11 if lvl == 5 else 12)
                shade(cell, "F2F2F2")
        row = table.add_row().cells
        lp = row[0].paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(lp, lgrp["runs"])
        ip = row[1].paragraphs[0]; ip.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(ip, igrp["runs"])

    set_col_widths(table, CELL_W)
    write_note_list(doc, f"Note al Libro {ROMAN[bi-1]}", footnotes)
    print(f"Libro {bi}: sez lat {len(lo)}/ita {len(io)} | note {len(footnotes)}")

def all_paragraphs(document):
    for p in document.paragraphs:
        yield p
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

for p in all_paragraphs(doc):
    for r in p.runs:
        force_font(r)

out = "/Users/macbook/Library/CloudStorage/Dropbox/ClaudeWS/Didattica/Anno accademico 2026_2027/LETTURA E CRITICA DEL TESTO FILOSOFICO MEDIEVALE/Contro_Accademici_latino_italiano.docx"
doc.save(out)
print("SAVED", out)
