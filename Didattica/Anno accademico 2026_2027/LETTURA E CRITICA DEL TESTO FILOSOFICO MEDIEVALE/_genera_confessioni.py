# -*- coding: utf-8 -*-
import re, html, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/private/tmp/claude-501/-Users-macbook-Library-CloudStorage-Dropbox-ClaudeWS/745dcfcf-3c6a-479d-ae69-a50371e3f234/scratchpad"
ROMAN = ["I","II","III","IV","V","VI","VII","VIII","IX","X","XI","XII","XIII"]

def get_body(path):
    d = open(path, encoding="cp1252", errors="replace").read()
    lo = d.lower()
    a = lo.find("<body")
    a = d.find(">", a) + 1
    b = lo.find("</body>")
    return d[a:b]

BLOCK = re.compile(r"<(h3|h4|h5|p)\b[^>]*>(.*?)</\1>", re.I | re.S)
NOTE  = re.compile(r'<a\s+[^>]*href="[^"]*#N(\d+)"[^>]*>.*?</a>', re.I | re.S)
KEY   = re.compile(r'<a\s+name="C_(\d{3})_(\d{3})_(\d{3})"', re.I)

def clean_inline(s):
    # protect note references -> placeholder
    s = NOTE.sub(lambda m: "\x00N%s\x00" % m.group(1), s)
    s = re.sub(r"<br\s*/?>", " ", s, flags=re.I)
    # drop tags we don't need, keep inner text
    s = re.sub(r"</?(?:a|u|sup|sub|font|span|h[1-6]|p|div)\b[^>]*>", "", s, flags=re.I)
    return s

def to_runs(s):
    """Return list of ('t', text, bold, italic) / ('n', number)."""
    s = clean_inline(s)
    parts = re.split(r"(<i>|</i>|<b>|</b>|\x00N\d+\x00)", s, flags=re.I)
    bold = ital = 0
    out = []
    for p in parts:
        if not p:
            continue
        pl = p.lower()
        if pl == "<i>": ital += 1; continue
        if pl == "</i>": ital = max(0, ital - 1); continue
        if pl == "<b>": bold += 1; continue
        if pl == "</b>": bold = max(0, bold - 1); continue
        m = re.match(r"\x00N(\d+)\x00", p)
        if m:
            out.append(("n", m.group(1)))
            continue
        txt = re.sub(r"<[^>]+>", "", p)
        txt = html.unescape(txt).replace("\xa0", " ")
        txt = re.sub(r"[ \t\r\n]+", " ", txt)
        if txt:
            out.append(("t", txt, bool(bold), bool(ital)))
    # trim leading/trailing whitespace of the whole run sequence
    ti = [i for i, r in enumerate(out) if r[0] == "t"]
    if ti:
        f = ti[0]; out[f] = ("t", out[f][1].lstrip(), out[f][2], out[f][3])
        l = ti[-1]; out[l] = ("t", out[l][1].rstrip(), out[l][2], out[l][3])
        out = [r for r in out if r[0] != "t" or r[1]]
    return out

def parse_book(path, lang):
    """-> ordered items + leading <p> lines.
    items: ('h', level, runs) or ('p', key, runs)
    Skips the book-header material (handled separately).
    lang 'lat': header is <h3> + first CAPS <h4>.
    lang 'ita': header is the leading <p> lines before the first <h4>."""
    body = get_body(path)
    raw = [(t.lower(), inner) for t, inner in BLOCK.findall(body)]
    items = []
    seen_first_h4 = False
    seen_first_key = False
    lead_p = []          # leading <p> before first h4 (Italian header lines)
    for tag, inner in raw:
        text_only = html.unescape(re.sub(r"<[^>]+>", "", inner)).replace("\xa0", " ").strip()
        if tag == "h3":
            continue
        if tag == "h4":
            if lang == "lat" and not seen_first_h4:
                seen_first_h4 = True      # CAPS theme line -> skip (latin only)
                continue
            seen_first_h4 = True
            if not text_only:
                continue
            items.append(("h", 4, to_runs(inner)))
            continue
        if tag == "h5":
            if not text_only:
                continue
            items.append(("h", 5, to_runs(inner)))
            continue
        # tag == p
        km = KEY.search(inner)
        if km:
            seen_first_key = True
            key = (km.group(2), km.group(3))
            items.append(("p", key, to_runs(inner)))
            continue
        if not seen_first_h4 and not seen_first_key:
            if text_only:
                lead_p.append(text_only)
            continue
        # non-keyed paragraph inside the body (spacers or stray)
        if text_only:
            items.append(("p", None, to_runs(inner)))
    return items, lead_p

LAT_THEME_FIX = {
    "ASCENSIO AD VERITAM": "ASCENSIO AD VERITATEM",
    "POST DEUM QUAESITUM ED COGNITUM": "POST DEUM QUAESITUM ET COGNITUM",
}

def _latin_theme(path):
    body = get_body(path)
    m = re.search(r"<h3\b[^>]*>.*?</h3>(.*?)<h4\b[^>]*>(.*?)</h4>", body, re.I | re.S)
    if not m:
        return ""
    t = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", "", m.group(2)))).strip()
    return LAT_THEME_FIX.get(t, t)

def parse_notes(path):
    d = open(path, encoding="cp1252", errors="replace").read()
    lo = d.lower(); a = lo.find("<body"); a = d.find(">", a) + 1; b = lo.find("</body>")
    body = d[a:b]
    notes = {}
    for m in re.finditer(r'<a\s+name="N(\d+)"[^>]*>.*?</a>(.*?)(?=<a\s+name="N\d+"|</body>|$)',
                         body, re.I | re.S):
        num = int(m.group(1))
        txt = m.group(2)
        txt = re.sub(r"<[^>]+>", "", txt)
        txt = html.unescape(txt).replace("\xa0", " ")
        txt = txt.replace("\r", " ").replace("\n", " ")
        txt = re.sub(r"\s+", " ", txt).strip()
        txt = re.sub(r"^[-–—\s]+", "", txt)
        notes[num] = txt
    return notes

# ---------- docx helpers ----------
def add_runs(par, runs, base_size=None):
    for r in runs:
        if r[0] == "n":
            run = par.add_run(r[1])
            run.font.superscript = True
            if base_size: run.font.size = base_size
        else:
            _, txt, b, i = r
            run = par.add_run(txt)
            run.bold = b
            run.italic = i
            if base_size: run.font.size = base_size

def set_col_widths(table, w):
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = w

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)

# ---------- build ----------
from docx.enum.section import WD_ORIENT

FONT = "Garamond"

def force_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rPr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)

doc = Document()
st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(12)                     # latino / italiano a corpo 12
st.paragraph_format.line_spacing = 1.15
for sname in ("Title", "Heading 1", "Heading 2", "Heading 3"):
    try:
        stx = doc.styles[sname]
        stx.font.name = FONT
        stx.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except KeyError:
        pass

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1.6)

CELL_W = Cm(13.3)

title = doc.add_heading("Agostino d’Ippona — Le Confessioni", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph("Testo latino e traduzione italiana a fronte")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

FONTE = (
    "Fonte di questa versione. Il testo latino e la traduzione italiana sono ripresi "
    "dall’edizione digitale della Nuova Biblioteca Agostiniana (Città Nuova Editrice), "
    "pubblicata sul sito www.augustinus.it:\n"
    "· testo latino: https://www.augustinus.it/latino/confessioni/ (conf_01–conf_13);\n"
    "· traduzione italiana: https://www.augustinus.it/italiano/confessioni/ (conf_01–conf_13).\n"
    "Consultazione: 7 settembre 2026. La numerazione di capitoli e paragrafi segue quella "
    "dell’edizione. Le note in calce a ciascun libro riproducono l’apparato di riferimenti "
    "(citazioni bibliche e classiche) della versione italiana; i rimandi nel testo latino "
    "sono stati omessi."
)
for line in FONTE.split("\n"):
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs: r.font.size = Pt(8.5); r.italic = True

for bi in range(1, 14):
    ii = "%02d" % bi
    lat_items, lat_lead = parse_book(f"{SRC}/latino_{ii}_libro.html", "lat")
    ita_items, ita_lead = parse_book(f"{SRC}/italiano_{ii}_libro.html", "ita")
    lat_notes = parse_notes(f"{SRC}/latino_{ii}_note.html")
    ita_notes = parse_notes(f"{SRC}/italiano_{ii}_note.html")
    parallel_notes = (lat_notes and ita_notes
                      and max(lat_notes) == max(ita_notes))

    # ---- book header ----
    doc.add_page_break()
    h = doc.add_heading(f"LIBRO {ROMAN[bi-1]}", level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lat_theme = _latin_theme(f"{SRC}/latino_{ii}_libro.html")
    ita_theme = ita_lead[1] if len(ita_lead) > 1 else ""
    hp = doc.add_paragraph(); hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = hp.add_run(lat_theme); r1.italic = True; r1.bold = True
    if ita_theme:
        hp2 = doc.add_paragraph(); hp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = hp2.add_run(ita_theme); r2.bold = True
        r2.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # ---- split each language into (heading-run, para) groups keyed by section ----
    def grouped(items):
        groups = {}       # key -> {'heads':[...], 'runs':[...]}
        order = []
        pending = []
        for it in items:
            if it[0] == "h":
                pending.append(it)
            else:  # p
                key = it[1]
                if key is None:
                    # attach stray paragraph to previous group if any, else skip
                    if order:
                        groups[order[-1]]["runs"].append(("t", " ", False, False))
                        groups[order[-1]]["runs"].extend(it[2])
                    continue
                if key not in groups:
                    groups[key] = {"heads": [], "runs": []}
                    order.append(key)
                groups[key]["heads"].extend(pending)
                groups[key]["runs"].extend(it[2])
                pending = []
        return order, groups

    lo, lg = grouped(lat_items)
    io, ig = grouped(ita_items)

    keys = list(lo)
    for k in io:
        if k not in lg:
            keys.append(k)

    # ---- table ----
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for c, lbl in zip(hdr, ("LATINO", "ITALIANO")):
        c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(lbl); rr.bold = True
        shade(c, "E8E8E8")

    for k in keys:
        lgrp = lg.get(k, {"heads": [], "runs": []})
        igrp = ig.get(k, {"heads": [], "runs": []})
        lh, ih = lgrp["heads"], igrp["heads"]
        nh = max(len(lh), len(ih))
        for hidx in range(nh):
            row = table.add_row().cells
            for cell, hs in ((row[0], lh), (row[1], ih)):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if hidx < len(hs):
                    lvl = hs[hidx][1]
                    add_runs(p, hs[hidx][2])
                    for rn in p.runs:
                        rn.bold = True
                        rn.italic = (lvl == 5)
                        rn.font.size = Pt(11 if lvl == 5 else 12)
                shade(cell, "F2F2F2")
        row = table.add_row().cells
        lp = row[0].paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(lp, [r for r in lgrp["runs"] if r[0] != "n"])   # no note markers in Latin
        ip = row[1].paragraphs[0]; ip.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(ip, igrp["runs"])

    set_col_widths(table, CELL_W)

    # ---- notes for this book ----
    def used_nums(order, groups):
        u = set()
        for k in order:
            g = groups.get(k)
            if not g: continue
            for r in g["runs"]:
                if r[0] == "n": u.add(int(r[1]))
        return u
    ita_used = used_nums(keys, ig)

    def write_note_list(heading, nums, table_src):
        doc.add_paragraph()
        np = doc.add_paragraph()
        nr = np.add_run(heading); nr.bold = True; nr.font.size = Pt(11)
        for num in sorted(nums):
            txt = table_src.get(num, "")
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            rn = p.add_run(f"{num}. "); rn.bold = True; rn.font.size = Pt(8.5)
            rt = p.add_run(txt); rt.font.size = Pt(8.5)

    write_note_list(f"Note al Libro {ROMAN[bi-1]}", ita_used, ita_notes)

    print(f"Libro {bi}: sez lat {len(lo)}/ita {len(io)} | note ita usate {len(ita_used)}/{max(ita_notes) if ita_notes else 0}")

# ---- final pass: force Garamond on every run (body + tables) ----
def all_paragraphs(document):
    for p in document.paragraphs:
        yield p
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

for p in all_paragraphs(doc):
    for r in p.runs:
        force_font(r)

out = "/Users/macbook/Library/CloudStorage/Dropbox/ClaudeWS/Didattica/Anno accademico 2026_2027/LETTURA E CRITICA DEL TESTO FILOSOFICO MEDIEVALE/Confessioni_latino_italiano.docx"
doc.save(out)
print("SAVED", out)
